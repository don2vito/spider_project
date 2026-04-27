# -*- coding: utf-8 -*-
"""
Weather History Scraper - Flask Proxy Server
Scrapes historical weather data from tianqihoubao.com
"""

import os
import re
import json
import time
import random
import logging
import threading
from datetime import datetime
from io import BytesIO

import requests
from flask import Flask, jsonify, send_file, request as flask_request
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# ============================================================
# Configuration
# ============================================================
BASE_URL = "https://www.tianqihoubao.com"
CACHE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cache")
DATA_START_YEAR = 2011

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
    "Accept-Encoding": "gzip, deflate",
    "Referer": "https://www.tianqihoubao.com/",
    "Connection": "keep-alive",
}

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)

app = Flask(__name__)

# Global rate limiter
_request_lock = threading.Lock()
_last_request_time = 0.0

# ============================================================
# HTTP Request Layer
# ============================================================
def fetch_page(url, retries=3):
    """Fetch a page with rate limiting, encoding handling, and retry logic."""
    global _last_request_time
    for attempt in range(retries):
        try:
            with _request_lock:
                now = time.time()
                elapsed = now - _last_request_time
                if elapsed < 2.0:
                    time.sleep(2.0 - elapsed)
                time.sleep(random.uniform(1.0, 3.0))
                _last_request_time = time.time()

                response = requests.get(url, headers=HEADERS, timeout=15)

            if response.status_code == 403:
                logger.warning(f"403 Forbidden for {url}, attempt {attempt + 1}/{retries}")
                time.sleep(5)
                continue

            response.raise_for_status()

            # Decode GBK content
            raw = response.content
            for encoding in ["gbk", "gb18030"]:
                try:
                    text = raw.decode(encoding)
                    return BeautifulSoup(text, "html.parser")
                except (UnicodeDecodeError, LookupError):
                    continue
            text = raw.decode("utf-8", errors="replace")
            return BeautifulSoup(text, "html.parser")

        except requests.RequestException as e:
            logger.error(f"Request failed for {url}: {e}, attempt {attempt + 1}/{retries}")
            if attempt < retries - 1:
                time.sleep(3)
    return None


# ============================================================
# Cache Layer
# ============================================================
def _ensure_cache_dir():
    os.makedirs(CACHE_DIR, exist_ok=True)


def _cache_path(key):
    return os.path.join(CACHE_DIR, f"{key}.json")


def get_cache(key, ttl=86400):
    """Read cache if exists and not expired."""
    fp = _cache_path(key)
    if not os.path.exists(fp):
        return None
    if time.time() - os.path.getmtime(fp) > ttl:
        return None
    try:
        with open(fp, "r", encoding="utf-8") as f:
            return json.load(f)
    except (json.JSONDecodeError, IOError):
        return None


def set_cache(key, data):
    """Write data to cache."""
    _ensure_cache_dir()
    fp = _cache_path(key)
    try:
        with open(fp, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except IOError as e:
        logger.error(f"Failed to write cache {key}: {e}")


# ============================================================
# Data Parsing Layer
# ============================================================
def parse_provinces():
    """Parse province list from /lishi/index.htm."""
    cache_key = "provinces"
    cached = get_cache(cache_key, ttl=7 * 86400)
    if cached:
        return cached

    soup = fetch_page(f"{BASE_URL}/lishi/index.htm")
    if not soup:
        return []

    provinces = []
    # The page uses: .citychk > dl > dt > a (province link with .htm extension)
    for a_tag in soup.select(".citychk dt a"):
        name = a_tag.get_text(strip=True)
        href = a_tag.get("href", "")
        # Extract abbr from href like "/lishi/hebei.htm"
        match = re.search(r"/lishi/([^/]+)\.htm", href)
        if match:
            abbr = match.group(1)
            provinces.append({"name": name, "abbr": abbr})

    set_cache(cache_key, provinces)
    logger.info(f"Parsed {len(provinces)} provinces")
    return provinces


def parse_cities(province_abbr):
    """Parse city list for a province from /lishi/{abbr}.htm."""
    cache_key = f"cities_{province_abbr}"
    cached = get_cache(cache_key, ttl=7 * 86400)
    if cached:
        return cached

    soup = fetch_page(f"{BASE_URL}/lishi/{province_abbr}.htm")
    if not soup:
        return []

    cities = []
    # The province page uses: .citychk > dl > dt > a (main city, .html)
    # and .citychk > dl > dd > a (district/county cities, .html)
    citychk = soup.select_one(".citychk")
    if not citychk:
        # Fallback: try #content
        citychk = soup.select_one("#content")
    if not citychk:
        return []

    for dl in citychk.select("dl"):
        # DT contains the main city name (prefecture-level)
        dt = dl.find("dt")
        if dt:
            a_tag = dt.find("a")
            if a_tag:
                name = a_tag.get_text(strip=True)
                href = a_tag.get("href", "")
                match = re.search(r"/lishi/([^/]+)\.html", href)
                if match:
                    pinyin = match.group(1)
                    cities.append({"name": name, "pinyin": pinyin, "is_main": True})

        # DD contains district/county cities
        for dd in dl.find_all("dd"):
            for a_tag in dd.find_all("a"):
                name = a_tag.get_text(strip=True)
                href = a_tag.get("href", "")
                match = re.search(r"/lishi/([^/]+)\.html", href)
                if match:
                    pinyin = match.group(1)
                    # Avoid duplicating the main city if it appears in DD too
                    if not any(c["pinyin"] == pinyin for c in cities):
                        cities.append({"name": name, "pinyin": pinyin, "is_main": False})

    set_cache(cache_key, cities)
    logger.info(f"Parsed {len(cities)} cities for province {province_abbr}")
    return cities


def parse_weather(city_pinyin, year, month):
    """Parse monthly weather data from /lishi/{city}/month/{YYYYMM}.html."""
    ym = f"{year}{month:02d}"
    cache_key = f"weather_{city_pinyin}_{ym}"
    cached = get_cache(cache_key, ttl=30 * 86400)
    if cached is not None:
        return cached

    # Skip future months
    now = datetime.now()
    if year > now.year or (year == now.year and month > now.month):
        return []

    url = f"{BASE_URL}/lishi/{city_pinyin}/month/{ym}.html"
    soup = fetch_page(url)
    if not soup:
        return []

    records = []
    # The actual table uses class="weather-table" with thead/tbody
    table = soup.select_one("table.weather-table")
    if not table:
        # Fallback: try class="b"
        table = soup.select_one("table.b")
    if not table:
        return []

    rows = table.find_all("tr")
    temp_pattern = re.compile(r"(-?\d+)℃")

    for row in rows:
        # Skip header rows (those inside thead)
        if row.find_parent("thead"):
            continue

        tds = row.find_all("td")
        if len(tds) < 4:
            continue

        # Date: link text like "2026年01月01日" or plain text
        date_link = tds[0].find("a")
        if date_link:
            date_text = date_link.get_text(strip=True)
        else:
            date_text = tds[0].get_text(strip=True)

        # Try to extract full date
        date_match = re.search(r"(\d{4})年(\d{1,2})月(\d{1,2})日", date_text)
        if date_match:
            date_str = f"{date_match.group(1)}-{int(date_match.group(2)):02d}-{int(date_match.group(3)):02d}"
        else:
            day_match = re.search(r"(\d{1,2})日", date_text)
            if day_match:
                date_str = f"{year}-{month:02d}-{int(day_match.group(1)):02d}"
            else:
                continue

        # Weather: " 晴 /  晴  " -> split by "/"
        weather_text = tds[1].get_text(strip=True)
        weather_text = re.sub(r"\s+", " ", weather_text).strip()
        if "/" in weather_text:
            parts = [p.strip() for p in weather_text.split("/")]
            weather_day = parts[0] if len(parts) > 0 else ""
            weather_night = parts[1] if len(parts) > 1 else parts[0]
        else:
            weather_day = weather_text
            weather_night = weather_text

        # Temperature: may be in <span class="temp-high">2℃</span> / <span class="temp-low">-10℃</span>
        # or plain text "2℃ / -10℃"
        temp_high_span = tds[2].find("span", class_="temp-high")
        temp_low_span = tds[2].find("span", class_="temp-low")
        if temp_high_span and temp_low_span:
            high_match = temp_pattern.search(temp_high_span.get_text())
            low_match = temp_pattern.search(temp_low_span.get_text())
            temp_high = int(high_match.group(1)) if high_match else None
            temp_low = int(low_match.group(1)) if low_match else None
        else:
            temp_text = tds[2].get_text(strip=True)
            temps = temp_pattern.findall(temp_text)
            if len(temps) >= 2:
                temp_high = int(temps[0])
                temp_low = int(temps[1])
            elif len(temps) == 1:
                temp_high = int(temps[0])
                temp_low = int(temps[0])
            else:
                temp_high = None
                temp_low = None

        # Wind: "东北风 1-3级 / 东北风 1-3级"
        wind_text = tds[3].get_text(strip=True)
        wind_text = re.sub(r"\s+", " ", wind_text).strip()
        if "/" in wind_text:
            wind_parts = [p.strip() for p in wind_text.split("/")]
            wind_day = wind_parts[0] if len(wind_parts) > 0 else ""
            wind_night = wind_parts[1] if len(wind_parts) > 1 else wind_parts[0]
        else:
            wind_day = wind_text
            wind_night = wind_text

        records.append({
            "date": date_str,
            "weather_day": weather_day,
            "weather_night": weather_night,
            "temp_high": temp_high,
            "temp_low": temp_low,
            "wind_day": wind_day,
            "wind_night": wind_night,
        })

    set_cache(cache_key, records)
    logger.info(f"Parsed {len(records)} records for {city_pinyin} {ym}")
    return records


# ============================================================
# Excel Export
# ============================================================
def create_excel(data, year, province_name, city_names):
    """Generate Excel file in memory."""
    wb = Workbook()
    ws = wb.active
    ws.title = "历史天气数据"

    # Styles
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1a365d", end_color="1a365d", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Microsoft YaHei", size=10)
    data_alignment = Alignment(horizontal="center", vertical="center")
    left_alignment = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="d0d0d0"),
        right=Side(style="thin", color="d0d0d0"),
        top=Side(style="thin", color="d0d0d0"),
        bottom=Side(style="thin", color="d0d0d0"),
    )

    even_fill = PatternFill(start_color="f7fafc", end_color="f7fafc", fill_type="solid")
    odd_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    neg_font = Font(name="Microsoft YaHei", size=10, color="FF0000")

    # Headers
    headers = ["一级地区", "二级地区", "日期", "白天天气", "夜间天气",
               "最高气温(℃)", "最低气温(℃)", "白天风力风向", "夜间风力风向"]
    col_widths = [12, 12, 14, 14, 14, 12, 12, 22, 22]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else "A" + chr(64 + col_idx - 26)].width = width

    # Data rows
    row_idx = 2
    for city_data in data:
        city_name = city_data.get("city", "")
        records = city_data.get("records", [])
        for rec in records:
            row_fill = even_fill if row_idx % 2 == 0 else odd_fill
            values = [
                province_name,
                city_name,
                rec.get("date", ""),
                rec.get("weather_day", ""),
                rec.get("weather_night", ""),
                rec.get("temp_high", ""),
                rec.get("temp_low", ""),
                rec.get("wind_day", ""),
                rec.get("wind_night", ""),
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                cell.fill = row_fill
                if col_idx == 3:  # Date column
                    cell.alignment = left_alignment
                else:
                    cell.alignment = data_alignment
                # Negative temperature in red
                if col_idx in (6, 7) and isinstance(val, (int, float)) and val < 0:
                    cell.font = neg_font
            row_idx += 1

    # Save to memory
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ============================================================
# Flask API Routes
# ============================================================
@app.route("/")
def index():
    return send_file("index.html")


@app.route("/api/provinces")
def api_provinces():
    provinces = parse_provinces()
    return jsonify({"provinces": provinces})


@app.route("/api/cities/<province_abbr>")
def api_cities(province_abbr):
    cities = parse_cities(province_abbr)
    return jsonify({"cities": cities})


@app.route("/api/weather")
def api_weather():
    year = flask_request.args.get("year", type=int)
    province = flask_request.args.get("province", "")
    cities_str = flask_request.args.get("cities", "")

    if not year or not province or not cities_str:
        return jsonify({"error": "Missing required parameters: year, province, cities"}), 400

    city_list = [c.strip() for c in cities_str.split(",") if c.strip()]
    if not city_list:
        return jsonify({"error": "No cities specified"}), 400

    all_data = []
    for city_pinyin in city_list:
        records = []
        for month in range(1, 13):
            month_records = parse_weather(city_pinyin, year, month)
            records.extend(month_records)
        # Sort by date
        records.sort(key=lambda x: x.get("date", ""))
        all_data.append({
            "city_pinyin": city_pinyin,
            "city": city_pinyin,  # Will be resolved to name later
            "records": records,
        })

    return jsonify({
        "province": province,
        "year": year,
        "data": all_data,
    })


@app.route("/api/weather/month")
def api_weather_month():
    """Fetch a single month's data (used for progressive loading)."""
    year = flask_request.args.get("year", type=int)
    month = flask_request.args.get("month", type=int)
    city = flask_request.args.get("city", "")

    if not all([year, month, city]):
        return jsonify({"error": "Missing parameters"}), 400

    records = parse_weather(city, year, month)
    return jsonify({
        "city_pinyin": city,
        "year": year,
        "month": month,
        "records": records,
    })


@app.route("/api/export")
def api_export():
    year = flask_request.args.get("year", type=int)
    province = flask_request.args.get("province", "")
    province_abbr = flask_request.args.get("province_abbr", "")
    cities_str = flask_request.args.get("cities", "")

    if not year or not province or not cities_str:
        return jsonify({"error": "Missing required parameters"}), 400

    city_list = [c.strip() for c in cities_str.split(",") if c.strip()]
    if not city_list:
        return jsonify({"error": "No cities specified"}), 400

    # Build city name mapping: load cities for this province (and all provinces if needed)
    city_name_map = {}  # pinyin -> name

    def _load_city_names_for_abbr(abbr):
        cache_file = os.path.join(CACHE_DIR, f"cities_{abbr}.json")
        if os.path.exists(cache_file):
            try:
                with open(cache_file, "r", encoding="utf-8") as f:
                    cities = json.load(f)
                for c in cities:
                    city_name_map[c["pinyin"]] = c["name"]
            except (json.JSONDecodeError, IOError):
                pass

    # First try the specified province
    if province_abbr:
        _load_city_names_for_abbr(province_abbr)

    # If some cities are still not resolved, search all cached province files
    unresolved = [c for c in city_list if c not in city_name_map]
    if unresolved and os.path.exists(CACHE_DIR):
        for fname in os.listdir(CACHE_DIR):
            if fname.startswith("cities_") and fname.endswith(".json"):
                if all(c in city_name_map for c in city_list):
                    break
                _load_city_names_for_abbr(fname.replace("cities_", "").replace(".json", ""))

    # Collect data
    all_data = []
    city_names = []
    for city_pinyin in city_list:
        records = []
        for month in range(1, 13):
            month_records = parse_weather(city_pinyin, year, month)
            records.extend(month_records)
        records.sort(key=lambda x: x.get("date", ""))

        city_name = city_name_map.get(city_pinyin, city_pinyin)
        all_data.append({"city": city_name, "records": records})
        city_names.append(city_name)

    # Generate Excel
    excel_buffer = create_excel(all_data, year, province, city_names)

    # Filename
    if len(city_names) == 1:
        filename = f"{year}_{province}_{city_names[0]}.xlsx"
    else:
        filename = f"{year}_{province}_{'_'.join(city_names[:3])}.xlsx"

    return send_file(
        excel_buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


# ============================================================
# Main
# ============================================================
if __name__ == "__main__":
    _ensure_cache_dir()
    logger.info("Starting Weather History Scraper server on http://127.0.0.1:5000")
    app.run(host="127.0.0.1", port=5000, debug=False)
